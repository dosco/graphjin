"""The whole-files design: keep each file whole, and give the agent tools to
find, search, read and query it. Nothing is split or rewritten.

Ingest writes, per corpus:
  text/<file>.txt   plain text with [page n] / [slide n] / [sheet name] markers
  tables.duckdb     every sheet and CSV as a raw grid and as a typed table
  catalog.json      title, description and structure of every file
"""

from __future__ import annotations

import csv
import json
import re
from datetime import date, datetime
from pathlib import Path

import duckdb
import openpyxl
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

MAX_READ_CHARS = 60_000


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _ident(s: str) -> str:
    s = re.sub(r"[^0-9a-zA-Z]+", "_", s.strip().lower()).strip("_")
    return s or "col"


def _sheet_rows(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _header_index(rows: list[list]) -> int | None:
    """The first row that looks like column titles: mostly text, with data below it."""
    for i, row in enumerate(rows[:30]):
        filled = [v for v in row if v not in (None, "")]
        if len(filled) < 2 or len(filled) < 0.6 * len([v for v in row if v is not None] or [0]):
            continue
        if all(isinstance(v, str) for v in filled) and i + 1 < len(rows) and any(v not in (None, "") for v in rows[i + 1]):
            return i
    return None


def _column_type(values: list) -> str:
    vals = [v for v in values if v not in (None, "")]
    if vals and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals):
        return "DOUBLE"
    if vals and all(isinstance(v, (datetime, date)) for v in vals):
        return "DATE"
    return "VARCHAR"


def _load_sheet(con, table: str, rows: list[list]) -> dict:
    width = max((len(r) for r in rows), default=0)
    letters = [openpyxl.utils.get_column_letter(i + 1) for i in range(width)]
    raw = f"{table}__raw"
    # Letters such as AS and BY are SQL keywords, so every column name is quoted.
    con.execute(f'create table "{raw}" (row_number INTEGER, {", ".join(f"{chr(34)}{c}{chr(34)} VARCHAR" for c in letters)})')
    con.executemany(f'insert into "{raw}" values ({", ".join(["?"] * (width + 1))})',
                    [[i + 1] + [_cell(v) if v is not None else None for v in (r + [None] * (width - len(r)))] for i, r in enumerate(rows)])
    info = {"raw_table": raw, "raw_columns": ["row_number"] + letters, "rows": len(rows)}
    h = _header_index(rows)
    if h is not None:
        names, seen = [], set()
        for j in range(width):
            v = rows[h][j] if j < len(rows[h]) else None
            n = _ident(str(v)) if v not in (None, "") else f"col_{letters[j].lower()}"
            while n in seen:
                n += "_2"
            seen.add(n)
            names.append(n)
        body = [r + [None] * (width - len(r)) for r in rows[h + 1:]]
        types = [_column_type([r[j] for r in body]) for j in range(width)]
        con.execute(f'create table "{table}" (source_row INTEGER, {", ".join(f"{chr(34)}{n}{chr(34)} {t}" for n, t in zip(names, types))})')

        def conv(v, t):
            if v in (None, ""):
                return None
            if t == "DATE":
                return v.date() if isinstance(v, datetime) else v
            if t == "DOUBLE":
                return float(v)
            return _cell(v)

        con.executemany(f'insert into "{table}" values ({", ".join(["?"] * (width + 1))})',
                        [[h + 2 + i] + [conv(v, t) for v, t in zip(r, types)] for i, r in enumerate(body) if any(v not in (None, "") for v in r)])
        info.update({"table": table, "header_row": h + 1, "columns": [f"{n} {t}" for n, t in zip(names, types)],
                     "table_rows": con.execute(f'select count(*) from "{table}"').fetchone()[0]})
    return info


def render(path: Path, con, stem: str) -> tuple[str, dict]:
    """Plain text for reading and searching, plus structure for the catalog."""
    ext = path.suffix.lower()
    if ext in (".md", ".txt", ".markdown"):
        return path.read_text(errors="replace"), {"kind": "text"}
    if ext == ".csv":
        text = path.read_text(errors="replace")
        rows = list(csv.reader(text.splitlines()))
        info = _load_sheet(con, stem, [[_num(v) for v in r] for r in rows])
        return text, {"kind": "table", "sheets": [{"sheet": path.name, **info}]}
    if ext == ".pdf":
        pages = [p.extract_text() or "" for p in PdfReader(path).pages]
        return "\n".join(f"[page {i}]\n{t}" for i, t in enumerate(pages, 1)), {"kind": "pdf", "pages": len(pages)}
    if ext == ".docx":
        doc = Document(path)
        out = []
        for block in doc.element.body.iterchildren():
            tag = block.tag.rsplit("}", 1)[-1]
            if tag == "p":
                t = "".join(n.text or "" for n in block.iter() if n.tag.endswith("}t"))
                if t.strip():
                    out.append(t)
            elif tag == "tbl":
                for tr in block.iter():
                    if tr.tag.endswith("}tr"):
                        cells = ["".join(n.text or "" for n in tc.iter() if n.tag.endswith("}t")) for tc in tr if tc.tag.endswith("}tc")]
                        out.append(" | ".join(cells))
        return "\n".join(out), {"kind": "document"}
    if ext == ".pptx":
        out = []
        for i, slide in enumerate(Presentation(path).slides, 1):
            out.append(f"[slide {i}]")
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    out.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        out.append(" | ".join(c.text for c in row.cells))
        return "\n".join(out), {"kind": "slides", "slides": sum(1 for line in out if line.startswith("[slide "))}
    if ext == ".xlsx":
        wb = openpyxl.load_workbook(path, data_only=True)
        out, sheets = [], []
        for ws in wb.worksheets:
            rows = _sheet_rows(ws)
            hidden = [k for k, d in ws.column_dimensions.items() if d.hidden]
            out.append(f"[sheet {ws.title}]" + (f" (hidden columns: {', '.join(hidden)})" if hidden else ""))
            for i, r in enumerate(rows, 1):
                if any(v not in (None, "") for v in r):
                    out.append(f"r{i}: " + " | ".join(_cell(v) for v in r))
            sheets.append({"sheet": ws.title, "hidden_columns": hidden,
                           **_load_sheet(con, f"{stem}__{_ident(ws.title)}", rows)})
        return "\n".join(out), {"kind": "spreadsheet", "sheets": sheets}
    raise ValueError(f"unsupported file {path.name}")


def _num(v: str):
    try:
        return float(v) if re.fullmatch(r"-?\d+(\.\d+)?", v.strip()) else v
    except ValueError:
        return v


DESCRIBE = """Describe this business file for a catalog that helps an assistant decide whether to open it.
Write two or three plain sentences: what the file is, what period or scope it covers, and what questions it can answer.
Mention if it supersedes or conflicts with another version when the file says so. No preamble.

File name: {name}
Content (may be cut off):
{text}"""


def ingest(corpus_dir: Path, out_dir: Path, gemini) -> dict:
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "text").mkdir(exist_ok=True)
    db = out_dir / "tables.duckdb"
    db.unlink(missing_ok=True)
    con = duckdb.connect(str(db))
    catalog, usage = [], {"promptTokenCount": 0, "candidatesTokenCount": 0}
    for path in sorted((corpus_dir / "files").iterdir()):
        stem = _ident(path.stem)
        text, info = render(path, con, stem)
        (out_dir / "text" / f"{path.name}.txt").write_text(text)
        description, u = gemini.text(DESCRIBE.format(name=path.name, text=text[:12_000]))
        for k in usage:
            usage[k] += u.get(k, 0)
        catalog.append({"file": path.name, "description": description.strip(), "characters": len(text), **info})
    con.close()
    (out_dir / "catalog.json").write_text(json.dumps(catalog, indent=1, default=str))
    return {"files": len(catalog), "usage": usage}


class FilesTools:
    """The four tools, over one ingested corpus."""

    def __init__(self, out_dir: Path):
        self.catalog = json.loads((out_dir / "catalog.json").read_text())
        self.text = {c["file"]: (out_dir / "text" / f"{c['file']}.txt").read_text() for c in self.catalog}
        self.con = duckdb.connect(str(out_dir / "tables.duckdb"), read_only=True)

    declarations = [
        {"name": "list_files", "description": "List every file in the company library with a short description and its structure: pages, slides, sheets, and the SQL tables loaded from each spreadsheet or CSV with their columns.",
         "parameters": {"type": "object", "properties": {}}},
        {"name": "search_files", "description": "Full-text search inside every file. Returns the files that match, each with matching lines and their page, slide, sheet row or line location.",
         "parameters": {"type": "object", "properties": {"query": {"type": "string", "description": "Words to look for"}}, "required": ["query"]}},
        {"name": "read_file", "description": "Read a file's full text. Pages, slides and sheets are marked like [page 3], [slide 2], [sheet Name] with rows as r12. For long files pass start and end to read part: page or slide numbers for PDFs and decks, line numbers otherwise.",
         "parameters": {"type": "object", "properties": {"file": {"type": "string"}, "start": {"type": "integer"}, "end": {"type": "integer"}}, "required": ["file"]}},
        {"name": "sql", "description": "Run a read-only DuckDB SQL query over the tables loaded from spreadsheets and CSV files (see list_files for table names and columns). Each sheet has a typed table using its header row and a <table>__raw table with every cell as text by column letter, with row_number. Use the raw table when a sheet has several tables, title rows or notes. Use SQL for counts, sums, filters and lookups so the numbers are exact.",
         "parameters": {"type": "object", "properties": {"query": {"type": "string"}}, "required": ["query"]}},
    ]

    def call(self, name: str, args: dict) -> str:
        try:
            return getattr(self, f"_{name}")(**args)
        except Exception as e:  # tool errors go back to the model
            return json.dumps({"error": f"{type(e).__name__}: {e}"})

    def _list_files(self) -> str:
        return json.dumps(self.catalog, default=str)

    def _search_files(self, query: str) -> str:
        terms = [t for t in re.findall(r"[\w\-.]+", query.lower()) if len(t) > 1]
        results = []
        for name, text in self.text.items():
            where, hits = "start", []
            score = 0
            for n, line in enumerate(text.splitlines(), 1):
                m = re.match(r"\[(page|slide) (\d+)\]|\[sheet (.+?)\]", line)
                if m:
                    where = f"{m.group(1)} {m.group(2)}" if m.group(1) else f"sheet {m.group(3)}"
                low = line.lower()
                matched = sum(1 for t in terms if t in low)
                if matched:
                    score += matched
                    if len(hits) < 12:
                        hits.append({"location": where, "line": n, "text": line[:300]})
            if score:
                results.append({"file": name, "score": score, "matches": hits})
        results.sort(key=lambda r: -r["score"])
        return json.dumps({"results": results[:8]})

    def _read_file(self, file: str, start: int | None = None, end: int | None = None) -> str:
        text = self.text.get(file)
        if text is None:
            return json.dumps({"error": f"no file named {file}"})
        paged = re.split(r"(?m)^(?=\[(?:page|slide) \d+\])", text)
        paged = [p for p in paged if p.strip()]
        if start is not None and len(paged) > 1:
            part = "".join(paged[max(start, 1) - 1:(end or start)])
        elif start is not None:
            lines = text.splitlines()
            part = "\n".join(lines[max(start, 1) - 1:(end or len(lines))])
        else:
            part = text
        note = ""
        if len(part) > MAX_READ_CHARS:
            part = part[:MAX_READ_CHARS]
            note = f"Cut at {MAX_READ_CHARS} characters; read the rest with start and end."
        return json.dumps({"file": file, "text": part, "note": note})

    def _sql(self, query: str) -> str:
        cur = self.con.execute(query)
        cols = [d[0] for d in cur.description]
        rows = cur.fetchmany(201)
        more = len(rows) > 200
        return json.dumps({"columns": cols, "rows": rows[:200], "truncated": more}, default=str)
