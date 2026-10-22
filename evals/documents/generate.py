"""Generate the synthetic document corpus and its questions.

Every file is built from one seeded data set, and every expected answer is
computed from that same data, so the answers are exact by construction.
Output: corpus/synthetic/files/* and corpus/synthetic/questions.yaml.
"""

from __future__ import annotations

import csv
import random
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

import yaml
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

SEED = 20261022
ROOT = Path(__file__).parent / "corpus" / "synthetic"
FILES = ROOT / "files"

CATEGORIES = ["Sensors", "Controllers", "Cables", "Enclosures", "Power"]
REGIONS = ["North", "South", "East", "West"]
WAREHOUSES = ["Rotterdam", "Leipzig", "Lyon"]


@dataclass
class Product:
    sku: str
    name: str
    category: str
    q3_price: float
    q4_price: float


@dataclass
class Order:
    region: str
    order_no: str
    day: date
    sku: str
    qty: int
    unit_price: float
    cost: float

    @property
    def net(self) -> float:
        return round(self.qty * self.unit_price, 2)


def build_data(rng: random.Random):
    nouns = {
        "Sensors": ["Pressure Sensor", "Thermal Probe", "Flow Meter", "Level Sensor"],
        "Controllers": ["PLC Module", "Loop Controller", "Motion Drive", "I/O Gateway"],
        "Cables": ["Shielded Cable", "Fiber Patch Lead", "Sensor Cable", "Power Harness"],
        "Enclosures": ["Steel Cabinet", "Polycarbonate Box", "DIN Enclosure", "Wall Housing"],
        "Power": ["DC Supply", "UPS Module", "Transformer", "Surge Guard"],
    }
    products: list[Product] = []
    n = 1001
    for cat in CATEGORIES:
        for noun in nouns[cat]:
            for variant in ("S", "X"):
                q3 = round(rng.uniform(18, 900), 2)
                q4 = q3
                products.append(Product(f"HI-{n}", f"{noun} {variant}{rng.randint(1, 9)}", cat, q3, q4))
                n += 1
    # Q4 repriced ten products; the CSV keeps Q3 prices.
    for p in rng.sample(products, 10):
        p.q4_price = round(p.q3_price * rng.choice([1.04, 1.06, 1.08, 0.95]), 2)

    orders: list[Order] = []
    seq = 1
    for region in REGIONS:
        for _ in range(rng.randint(22, 30)):
            p = rng.choice(products)
            day = date(2025, 1, 1) + timedelta(days=rng.randint(0, 364))
            price = p.q4_price if day >= date(2025, 10, 1) else p.q3_price
            orders.append(Order(region, f"SO-2025-{seq:04d}", day, p.sku, rng.randint(1, 40), price, round(price * 0.61, 2)))
            seq += 1
    cancelled = rng.choice([o for o in orders if o.region == "South"]).order_no

    returns = []
    for o in rng.sample(orders, 12):
        if o.order_no == cancelled:
            continue
        returns.append({"order_no": o.order_no, "sku": o.sku, "qty": rng.randint(1, max(1, o.qty // 3)),
                        "reason": rng.choice(["Damaged in transit", "Wrong item", "Customer cancelled", "Defective"])})

    inventory = []
    for i in range(1500):
        p = rng.choice(products)
        wh = rng.choice(WAREHOUSES)
        on_hand = rng.randint(0, 400)
        inventory.append({"warehouse": wh, "sku": p.sku, "name": p.name, "category": p.category,
                          "bin": f"{wh[:3].upper()}-{rng.randint(1, 40):02d}-{rng.choice('ABCDEF')}",
                          "on_hand": on_hand, "reserved": rng.randint(0, on_hand),
                          "counted": date(2025, 12, 1) + timedelta(days=rng.randint(0, 20))})
    return products, orders, cancelled, returns, inventory


def write_sales(products, orders, cancelled, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales 2025"
    ws["A1"] = "Halden Instruments — Regional Sales 2025"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:G1")
    ws["A2"] = "Confidential. Prepared by Finance. All amounts in EUR."
    ws.merge_cells("A2:G2")
    row = 4
    for region in REGIONS:
        ws.cell(row=row, column=1, value=f"Region: {region}").font = Font(bold=True)
        row += 1
        for col, title in enumerate(["Order Date", "Order No", "SKU", "Qty", "Unit Price", "Net Amount", "Internal Cost"], 1):
            ws.cell(row=row, column=col, value=title).font = Font(bold=True)
        row += 1
        region_orders = sorted((o for o in orders if o.region == region), key=lambda o: o.day)
        for o in region_orders:
            ws.cell(row=row, column=1, value=o.day).number_format = "DD.MM.YYYY"
            ws.cell(row=row, column=2, value=o.order_no)
            ws.cell(row=row, column=3, value=o.sku)
            ws.cell(row=row, column=4, value=o.qty)
            ws.cell(row=row, column=5, value=o.unit_price).number_format = "#,##0.00"
            ws.cell(row=row, column=6, value=o.net).number_format = "#,##0.00"
            ws.cell(row=row, column=7, value=round(o.cost * o.qty, 2))
            row += 1
        subtotal = round(sum(o.net for o in region_orders), 2)
        ws.cell(row=row, column=2, value=f"Subtotal {region}").font = Font(bold=True)
        ws.cell(row=row, column=6, value=subtotal).number_format = "#,##0.00"
        row += 3
    ws.cell(row=row, column=1, value=f"Notes: order {cancelled} was cancelled after booking; "
                                     "exclude it from all totals. Subtotals above still include it.")
    ws.column_dimensions["G"].hidden = True

    rs = wb.create_sheet("Returns")
    rs.append(["Returns logged in 2025"])
    rs.append([])
    rs.append(["Order No", "SKU", "Qty Returned", "Reason"])
    wb.save(path)


def write_returns(wb_path: Path, returns) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(wb_path)
    rs = wb["Returns"]
    for r in returns:
        rs.append([r["order_no"], r["sku"], r["qty"], r["reason"]])
    wb.save(wb_path)


def write_prices(products, csv_path: Path, xlsx_path: Path) -> None:
    with csv_path.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["SKU", "Description", "Category", "List Price (EUR)"])
        for p in products:
            w.writerow([p.sku, p.name, p.category, f"{p.q3_price:.2f}"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Q4 2025"
    ws["A1"] = "Price list effective 1 October 2025 — replaces the Q3 list"
    ws.merge_cells("A1:D1")
    ws.append([])
    ws.append(["SKU", "Description", "Category", "List Price (EUR)"])
    for p in products:
        ws.append([p.sku, p.name, p.category, p.q4_price])
    wb.save(xlsx_path)


def write_inventory(inventory, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(["Warehouse", "SKU", "Product", "Category", "Bin", "On Hand", "Reserved", "Last Count"])
    for r in inventory:
        ws.append([r["warehouse"], r["sku"], r["name"], r["category"], r["bin"], r["on_hand"], r["reserved"], r["counted"]])
    wb.save(path)


POLICIES = {
    2024: {"tier1": 62, "tier2": 48, "hotel_cap": 160, "approval": 2500, "class": "economy for flights under 6 hours"},
    2025: {"tier1": 70, "tier2": 52, "hotel_cap": 185, "approval": 3000, "class": "economy for flights under 8 hours"},
}


def write_policy(year: int, path: Path) -> None:
    p = POLICIES[year]
    doc = Document()
    doc.add_heading(f"Halden Instruments Travel and Expense Policy {year}", 0)
    doc.add_paragraph(f"Effective 1 January {year}. This policy replaces all earlier travel policies.")
    doc.add_heading("1. Booking", 1)
    doc.add_paragraph(f"Book flights in {p['class']}; business class needs CFO approval.")
    doc.add_heading("2. Per diem", 1)
    doc.add_paragraph("Per diem covers meals and incidentals and is paid per travel day.")
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate([("City tier", "Per diem (EUR)"), ("Tier 1 (capitals)", str(p["tier1"])), ("Tier 2 (all other cities)", str(p["tier2"]))]):
        table.cell(i, 0).text, table.cell(i, 1).text = k, v
    doc.add_heading("3. Hotels", 1)
    doc.add_paragraph(f"Hotel rates are capped at EUR {p['hotel_cap']} per night including taxes.")
    doc.add_heading("4. Approval", 1)
    doc.add_paragraph(f"Trips with an estimated total cost above EUR {p['approval']:,} need approval from a department head before booking.")
    doc.save(path)


ANNUAL = {
    "quarters": {"Q1": 4.82, "Q2": 5.11, "Q3": 5.37, "Q4": 5.94},
    "headcount": {"Rotterdam": 212, "Leipzig": 148, "Lyon": 97, "Remote": 41},
    "ceo": "Marta Vos",
    "plant": "Leipzig",
    "margin": 38.6,
}


def write_annual_report(path: Path) -> None:
    styles = getSampleStyleSheet()
    story = [Paragraph("Halden Instruments Annual Report 2025", styles["Title"]),
             Paragraph("Letter from the CEO", styles["Heading2"]),
             Paragraph(f"2025 was our strongest year. Revenue grew in every quarter, and our new {ANNUAL['plant']} "
                       "assembly line reached full output in September. Gross margin for the year was "
                       f"{ANNUAL['margin']}%. — {ANNUAL['ceo']}, Chief Executive Officer", styles["BodyText"]),
             Spacer(1, 12), Paragraph("Revenue by quarter (EUR millions)", styles["Heading2"])]
    data = [["Quarter", "Revenue"]] + [[q, f"{v:.2f}"] for q, v in ANNUAL["quarters"].items()]
    data.append(["Total", f"{sum(ANNUAL['quarters'].values()):.2f}"])
    t = Table(data)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
    story += [t, PageBreak(), Paragraph("People", styles["Heading2"]),
              Paragraph("Headcount at 31 December 2025, by site:", styles["BodyText"])]
    h = Table([["Site", "Employees"]] + [[k, str(v)] for k, v in ANNUAL["headcount"].items()])
    h.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey)]))
    story += [h, Spacer(1, 12),
              Paragraph("Outlook", styles["Heading2"]),
              Paragraph("For 2026 we expect revenue growth between 6% and 9%, led by the Controllers category.", styles["BodyText"])]
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


OPS = {"on_time": 94.2, "lines": {"Line A": 0.8, "Line B": 1.9, "Line C": 0.6}, "supplier": "Brandt Metallwerk"}


def write_ops_review(path: Path) -> None:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Operations Review Q3 2025"
    s.placeholders[1].text = "Halden Instruments — Manufacturing"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Delivery performance"
    s.placeholders[1].text = f"On-time delivery: {OPS['on_time']}% (target 95%)\nLate shipments mostly enclosure orders"
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Defect rate by line (%)"
    rows = [("Line", "Defect rate %")] + [(k, f"{v}") for k, v in OPS["lines"].items()]
    table = s.shapes.add_table(len(rows), 2, Inches(1), Inches(1.6), Inches(6), Inches(1.6)).table
    for i, (a, b) in enumerate(rows):
        table.cell(i, 0).text, table.cell(i, 1).text = a, b
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Supplier issues"
    s.placeholders[1].text = f"{OPS['supplier']} missed three steel deliveries; second source qualified in October"
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(20)
    prs.save(path)


SUPPLIERS = [("Brandt Metallwerk", "steel enclosures", date(2026, 3, 31), 90),
             ("Kowalski Electronics", "sensor boards", date(2026, 7, 15), 60),
             ("Nordkabel AB", "cables", date(2026, 1, 31), 30)]


def write_supplier_notes(path: Path) -> None:
    lines = ["# Supplier contracts", "", "Procurement notes, updated December 2025.", ""]
    for name, what, end, notice in SUPPLIERS:
        lines += [f"## {name}", f"- Supplies: {what}", f"- Contract ends: {end:%d %B %Y}",
                  f"- Notice period: {notice} days before the end date", ""]
    path.write_text("\n".join(lines))


def questions(products, orders, cancelled, returns, inventory):
    by_sku = {p.sku: p for p in products}
    qs = []

    def add(qid, kind, prompt, answer, files, **extra):
        qs.append({"id": qid, "kind": kind, "prompt": prompt, "answer": answer, "files": files, **extra})

    south = [o for o in orders if o.region == "South"]
    add("s01", "number", "What were total net sales in the South region in 2025, excluding cancelled orders? Answer in EUR.",
        round(sum(o.net for o in south if o.order_no != cancelled), 2), ["sales_2025.xlsx"], tolerance_pct=0.5)
    add("s02", "number", "How many sales orders did the North region book in 2025?",
        sum(1 for o in orders if o.region == "North"), ["sales_2025.xlsx"], tolerance_pct=0)
    biggest = max((o for o in orders if o.order_no != cancelled), key=lambda o: o.net)
    add("s03", "text", "Excluding cancelled orders, which order had the largest net amount in 2025? Give the order number.",
        biggest.order_no, ["sales_2025.xlsx"])
    q4 = [o for o in orders if o.day >= date(2025, 10, 1) and o.order_no != cancelled]
    add("s04", "number", "Across all regions, how many units were sold in Q4 2025 (October to December), excluding cancelled orders?",
        sum(o.qty for o in q4), ["sales_2025.xlsx"], tolerance_pct=0)
    add("s05", "number", "How many units were returned in total in 2025?",
        sum(r["qty"] for r in returns), ["sales_2025.xlsx"], tolerance_pct=0)
    top_reason = max({r["reason"] for r in returns}, key=lambda k: (sum(r["qty"] for r in returns if r["reason"] == k), k))
    add("s06", "text", "Which return reason accounts for the most returned units?", top_reason, ["sales_2025.xlsx"])

    repriced = [p for p in products if p.q4_price != p.q3_price]
    p0 = repriced[0]
    add("p01", "number", f"What is the current list price of SKU {p0.sku}? Answer in EUR.", p0.q4_price,
        ["price_list_q3_2025.csv", "price_list_q4_2025.xlsx"], tolerance_pct=0)
    add("p02", "number", "How many products changed price between the Q3 and Q4 2025 price lists?",
        len(repriced), ["price_list_q3_2025.csv", "price_list_q4_2025.xlsx"], tolerance_pct=0)
    cheapest = min((p for p in products if p.category == "Controllers"), key=lambda p: p.q4_price)
    add("p03", "text", "Which product in the Controllers category has the lowest current list price? Give its SKU.",
        cheapest.sku, ["price_list_q4_2025.xlsx"])

    sku_i = inventory[7]["sku"]
    avail = sum(r["on_hand"] - r["reserved"] for r in inventory if r["sku"] == sku_i and r["warehouse"] == "Leipzig")
    add("i01", "number", f"How many units of {sku_i} are available (on hand minus reserved) in the Leipzig warehouse?",
        avail, ["inventory_snapshot.xlsx"], tolerance_pct=0)
    add("i02", "number", "What is the total on-hand quantity of Cables across all warehouses?",
        sum(r["on_hand"] for r in inventory if r["category"] == "Cables"), ["inventory_snapshot.xlsx"], tolerance_pct=0)
    value = sum((r["on_hand"]) * by_sku[r["sku"]].q4_price for r in inventory if r["warehouse"] == "Rotterdam")
    add("i03", "number", "What is the total value of on-hand stock in Rotterdam at current list prices? Answer in EUR.",
        round(value, 2), ["inventory_snapshot.xlsx", "price_list_q4_2025.xlsx"], tolerance_pct=0.5)

    add("t01", "number", "What is the current per diem for a trip to a capital city? Answer in EUR.",
        POLICIES[2025]["tier1"], ["travel_policy_2024.docx", "travel_policy_2025.docx"], tolerance_pct=0)
    add("t02", "number", "Above what estimated trip cost does a trip need department head approval? Answer in EUR.",
        POLICIES[2025]["approval"], ["travel_policy_2025.docx"], tolerance_pct=0)
    add("t03", "number", "By how many euros per night did the hotel cap rise from the 2024 policy to the 2025 policy?",
        POLICIES[2025]["hotel_cap"] - POLICIES[2024]["hotel_cap"], ["travel_policy_2024.docx", "travel_policy_2025.docx"], tolerance_pct=0)

    add("a01", "number", "What was Halden Instruments' total revenue in 2025, in EUR millions?",
        round(sum(ANNUAL["quarters"].values()), 2), ["annual_report_2025.pdf"], tolerance_pct=0.5)
    add("a02", "number", "How many employees work at the Leipzig site?", ANNUAL["headcount"]["Leipzig"],
        ["annual_report_2025.pdf"], tolerance_pct=0)
    add("a03", "text", "Who is the Chief Executive Officer of Halden Instruments?", ANNUAL["ceo"], ["annual_report_2025.pdf"])

    add("o01", "number", "What was on-time delivery in Q3 2025, as a percentage?", OPS["on_time"],
        ["ops_review_q3_2025.pptx"], tolerance_pct=0)
    worst = max(OPS["lines"], key=OPS["lines"].get)
    add("o02", "text", "Which production line had the highest defect rate in Q3 2025?", worst, ["ops_review_q3_2025.pptx"])

    first = min(SUPPLIERS, key=lambda s: s[2] - timedelta(days=s[3]))
    add("m01", "text", "Which supplier contract has the earliest notice deadline (end date minus notice period)?", first[0],
        ["supplier_notes.md"])
    brandt = SUPPLIERS[0]
    add("m02", "text", "By what date must notice be given to end the Brandt Metallwerk contract? Give the date.",
        (brandt[2] - timedelta(days=brandt[3])).isoformat(), ["supplier_notes.md"], match="date")
    return qs


def main() -> None:
    rng = random.Random(SEED)
    FILES.mkdir(parents=True, exist_ok=True)
    products, orders, cancelled, returns, inventory = build_data(rng)
    write_sales(products, orders, cancelled, FILES / "sales_2025.xlsx")
    write_returns(FILES / "sales_2025.xlsx", returns)
    write_prices(products, FILES / "price_list_q3_2025.csv", FILES / "price_list_q4_2025.xlsx")
    write_inventory(inventory, FILES / "inventory_snapshot.xlsx")
    write_policy(2024, FILES / "travel_policy_2024.docx")
    write_policy(2025, FILES / "travel_policy_2025.docx")
    write_annual_report(FILES / "annual_report_2025.pdf")
    write_ops_review(FILES / "ops_review_q3_2025.pptx")
    write_supplier_notes(FILES / "supplier_notes.md")
    qs = questions(products, orders, cancelled, returns, inventory)
    (ROOT / "questions.yaml").write_text(yaml.safe_dump({"corpus": "synthetic", "seed": SEED, "questions": qs}, sort_keys=False, allow_unicode=True))
    print(f"wrote {len(list(FILES.iterdir()))} files and {len(qs)} questions to {ROOT}")


if __name__ == "__main__":
    main()
